"""Writes one Excel file per feature under test_data/<Feature>/ (or
test_data/WTA/<Feature>/ for any feature belonging to the WTA parent
module - mirrors tests/wta/ and pages/wta/, see AGENT.md), columns per the
spec: Test Case ID, Feature, Title, Description, Precondition, Test Data,
Steps, Expected Result, Actual Result, Assertions, Result."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config.settings import TEST_DATA_DIR, feature_dir

COLUMNS = ["Test Case ID", "Feature", "Title", "Description", "Precondition",
           "Test Data", "Steps", "Expected Result", "Actual Result",
           "Assertions", "Result"]

RESULT_FILL = {
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
    "BLOCKED": PatternFill("solid", fgColor="FFEB9C"),
}


def _assertions_text(assertions) -> str:
    lines = []
    for a in assertions:
        mark = "PASS" if a["passed"] else "FAIL"
        lines.append(f"[{mark}] {a['text']}")
    return "\n".join(lines) if lines else "-"


def write_feature_excel(feature: str, cases: list) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = feature[:31]

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")

    for case in cases:
        row = [
            case["id"], case["feature"], case["title"], case["description"],
            case["precondition"], case["test_data"], case["steps"],
            case["expected"], case["actual"], _assertions_text(case["assertions"]),
            case["result"],
        ]
        ws.append(row)
        result_cell = ws.cell(row=ws.max_row, column=len(COLUMNS))
        result_cell.fill = RESULT_FILL.get(case["result"], PatternFill())
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = [14, 12, 28, 34, 22, 22, 34, 30, 34, 30, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    out_dir = feature_dir(TEST_DATA_DIR, feature)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{feature}_Test_Cases.xlsx"
    wb.save(out_path)
    return out_path
